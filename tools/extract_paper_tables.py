"""Used once to freeze the tables as printed in the submitted manuscript (main text, .docx) and Supporting
Information (.docx) into reference/paper_tables/, so that pipeline/40_verify.py can compare regenerated tables with
the printed ones cell by cell. Header rows are dropped; only data rows are kept, as printed.

Usage: python tools/extract_paper_tables.py <manuscript.docx> <supporting_information.docx>"""
import sys, pathlib, csv
import docx
from docx.oxml.ns import qn

OUT = pathlib.Path(__file__).resolve().parents[1] / "reference" / "paper_tables"; OUT.mkdir(parents=True, exist_ok=True)


def cell_text(tc):
    return " ".join("".join(n.text or "" for n in p.iter() if n.tag in (qn("w:t"), qn("m:t")))
                    for p in tc.iter(qn("w:p"))).strip()


def rows_of(tbl):
    return [[cell_text(tc) for tc in tr.iter(qn("w:tc"))] for tr in tbl._tbl.iter(qn("w:tr"))]


def write(name, rows):
    with open(OUT / f"Table_{name}.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


ms = docx.Document(sys.argv[1]).tables
assert len(ms) == 8, len(ms)
# number of header rows of each main text table (I to VIII); Table I (literature) and VII (checklist) carry no result
spec = {"II": (1, 1), "III": (2, 2), "IV": (3, 1), "V": (4, 2), "VI": (5, 2), "VIII": (7, 1)}
for name, (i, nh) in spec.items():
    rows = rows_of(ms[i])[nh:]
    if name == "VIII":
        rows = [[r[0], r[2], r[3]] for r in rows]          # Group, Primary, Full (the definition column is text)
    write(name, rows)
si = docx.Document(sys.argv[2]).tables
assert len(si) == 7, len(si)
for k, nh in enumerate([1, 2, 2, 2, 2, 1, 1]):
    write(f"S{k + 1}", rows_of(si[k])[nh:])
print("written", sorted(p.name for p in OUT.glob("*.csv")))
